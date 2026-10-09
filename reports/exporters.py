"""Export feedback and study packs to DOCX or PDF."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.shared import Pt
from fpdf import FPDF


def _safe(text: str) -> str:
    return (text or "").replace("\x00", "").strip()


def _add_docx_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(_safe(text), level=level)


def _add_docx_para(doc: Document, label: str, body: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: ")
    run.bold = True
    run.font.size = Pt(11)
    body_run = p.add_run(_safe(body))
    body_run.font.size = Pt(11)


def write_feedback_docx(
    path: Path,
    *,
    candidate_name: str,
    candidate_email: str,
    feedback: str,
    qa_pairs: Iterable[tuple[str, str]],
) -> Path:
    doc = Document()
    _add_docx_heading(doc, "Interview Feedback Report")
    doc.add_paragraph(f"Candidate: {_safe(candidate_name)}")
    doc.add_paragraph(f"Email: {_safe(candidate_email)}")
    _add_docx_heading(doc, "Overall Feedback", level=2)
    doc.add_paragraph(_safe(feedback) or "No feedback available.")
    _add_docx_heading(doc, "Your Answers", level=2)
    for index, (question, answer) in enumerate(qa_pairs, start=1):
        _add_docx_heading(doc, f"Q{index}", level=3)
        _add_docx_para(doc, "Question", question)
        _add_docx_para(doc, "Your answer", answer)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def write_study_pack_docx(
    path: Path,
    *,
    candidate_name: str,
    candidate_email: str,
    items: list[dict[str, str]],
) -> Path:
    doc = Document()
    _add_docx_heading(doc, "Interview Prep Study Pack")
    doc.add_paragraph(f"Prepared for: {_safe(candidate_name)} ({_safe(candidate_email)})")
    doc.add_paragraph(
        "Use these expected questions, model answers, and concept notes to prepare for similar interviews."
    )
    for index, item in enumerate(items, start=1):
        _add_docx_heading(doc, f"Expected Question {index}", level=2)
        _add_docx_para(doc, "Question", item.get("question", ""))
        _add_docx_para(doc, "Solid answer", item.get("model_answer", ""))
        _add_docx_para(doc, "Concept explanation", item.get("concept_explanation", ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


class _ReportPDF(FPDF):
    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 8, f"Page {self.page_no()}", align="C")


def _pdf_write_wrapped(pdf: FPDF, text: str, size: int = 11) -> None:
    pdf.set_font("Helvetica", size=size)
    pdf.set_text_color(30, 40, 38)
    # fpdf2 latin-1 by default; sanitize for common unicode
    cleaned = (
        _safe(text)
        .replace("—", "-")
        .replace("–", "-")
        .replace("•", "-")
        .replace("“", '"')
        .replace("”", '"')
        .replace("’", "'")
        .replace("⭐", "*")
        .replace("✅", "[ok]")
        .replace("❌", "[x]")
        .replace("🟢", "[*]")
    )
    pdf.multi_cell(0, 6, cleaned)
    pdf.ln(2)


def write_feedback_pdf(
    path: Path,
    *,
    candidate_name: str,
    candidate_email: str,
    feedback: str,
    qa_pairs: Iterable[tuple[str, str]],
) -> Path:
    pdf = _ReportPDF()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Interview Feedback Report", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    _pdf_write_wrapped(pdf, f"Candidate: {candidate_name}")
    _pdf_write_wrapped(pdf, f"Email: {candidate_email}")
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Overall Feedback", new_x="LMARGIN", new_y="NEXT")
    _pdf_write_wrapped(pdf, feedback or "No feedback available.")
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Your Answers", new_x="LMARGIN", new_y="NEXT")
    for index, (question, answer) in enumerate(qa_pairs, start=1):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, f"Q{index}", new_x="LMARGIN", new_y="NEXT")
        _pdf_write_wrapped(pdf, f"Question: {question}")
        _pdf_write_wrapped(pdf, f"Your answer: {answer}")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


def write_study_pack_pdf(
    path: Path,
    *,
    candidate_name: str,
    candidate_email: str,
    items: list[dict[str, str]],
) -> Path:
    pdf = _ReportPDF()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Interview Prep Study Pack", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    _pdf_write_wrapped(pdf, f"Prepared for: {candidate_name} ({candidate_email})")
    _pdf_write_wrapped(
        pdf,
        "Use these expected questions, model answers, and concept notes to prepare for similar interviews.",
    )
    for index, item in enumerate(items, start=1):
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 8, f"Expected Question {index}", new_x="LMARGIN", new_y="NEXT")
        _pdf_write_wrapped(pdf, f"Question: {item.get('question', '')}")
        _pdf_write_wrapped(pdf, f"Solid answer: {item.get('model_answer', '')}")
        _pdf_write_wrapped(pdf, f"Concept explanation: {item.get('concept_explanation', '')}")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path
